"""Проверка структуры Excel, запись факта и ключевых событий в тот же файл."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

PLAN_SHEET = "FCF 2026 ПЛАН"
FACT_SHEET = "FCF 2026 ФАКТ"
EVENTS_SHEET = "Ключевые события"
BASE_YEAR = 2026

FACT_ALIASES = {
    "Ремонт квартиры": ("Ремонт квартиры", "Ремонт в квартире"),
}
PLAN_ALIASES = {
    "Ремонт квартиры": ("Ремонт квартиры", "Ремонт в квартире"),
}

EVENTS_HEADERS = ["Год", "Месяц", "Статья", "Событие", "Источник", "Ключ", "Скрыто"]


class ExcelStructureError(ValueError):
    """Файл есть, но листы/строки не совпали с ожидаемой моделью FCF."""


def _label(value) -> str:
    return " ".join(str(value or "").split())


def _matches(actual: str, expected: str, aliases: dict[str, tuple[str, ...]]) -> bool:
    names = aliases.get(expected, (expected,))
    return actual in names


def _seed():
    from . import seed_excel as s
    return s


def validate_workbook(wb, excel_name: str = "") -> list[str]:
    s = _seed()
    errors: list[str] = []
    if PLAN_SHEET not in wb.sheetnames:
        errors.append(f"нет листа «{PLAN_SHEET}»")
    if FACT_SHEET not in wb.sheetnames:
        errors.append(f"нет листа «{FACT_SHEET}»")
    if errors:
        return errors

    fact = wb[FACT_SHEET]
    plan = wb[PLAN_SHEET]
    year_cell = fact.cell(1, 3).value
    try:
        year_ok = int(float(year_cell)) == BASE_YEAR
    except (TypeError, ValueError):
        year_ok = False
    if not year_ok:
        errors.append(f"в «{FACT_SHEET}» C1 должно быть {BASE_YEAR}, сейчас {year_cell!r}")

    for month in range(1, 13):
        raw = fact.cell(2, 2 + month).value
        try:
            got = int(float(raw))
        except (TypeError, ValueError):
            got = None
        if got != month:
            errors.append(f"в факте строка 2, месяц {month}: ожидали {month}, получили {raw!r}")
            break

    for cat, row in {**s.FACT_INCOME, **s.FACT_EXPENSE}.items():
        actual = _label(fact.cell(row, 1).value)
        if not _matches(actual, cat, FACT_ALIASES):
            errors.append(f"факт строка {row}: ожидали «{cat}», получили «{actual or 'пусто'}»")

    for cat, row in {**s.PLAN_INCOME, **s.PLAN_EXPENSE}.items():
        actual = _label(plan.cell(row, 1).value)
        if not _matches(actual, cat, PLAN_ALIASES):
            errors.append(f"план строка {row}: ожидали «{cat}», получили «{actual or 'пусто'}»")

    if excel_name and "Бюджет 2026" not in excel_name:
        errors.append(f"имя файла без «Бюджет 2026»: {excel_name}")
    return errors


def validate_excel(path: Path | None = None) -> dict:
    s = _seed()
    excel = path or s.find_excel()
    if not excel or not excel.exists():
        raise FileNotFoundError("В папке проекта нет файла *Бюджет 2026.xlsx")
    try:
        wb = load_workbook(excel, data_only=False)
    except InvalidFileException as exc:
        raise ExcelStructureError(f"не открывается как Excel: {exc}") from exc
    try:
        errors = validate_workbook(wb, excel.name)
    finally:
        wb.close()
    if errors:
        raise ExcelStructureError("; ".join(errors[:8]) + ("…" if len(errors) > 8 else ""))
    return {"ok": True, "excel": excel.name, "path": str(excel), "errors": []}


def fact_row(category: str) -> int | None:
    s = _seed()
    return s.FACT_INCOME.get(category) or s.FACT_EXPENSE.get(category)


def _write_fact_cells(wb, cells: list[dict]) -> int:
    s = _seed()
    ws = wb[FACT_SHEET]
    written = 0
    for item in cells:
        cat = item["category"]
        month = int(item["month"])
        row = fact_row(cat)
        if not row or month < 1 or month > 12:
            continue
        ws.cell(row, 2 + month).value = s._num(item.get("value"))
        written += 1
    return written


def _write_events_sheet(wb, events: list[dict]) -> None:
    if EVENTS_SHEET in wb.sheetnames:
        ws = wb[EVENTS_SHEET]
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(EVENTS_SHEET)
    for col, head in enumerate(EVENTS_HEADERS, 1):
        ws.cell(1, col).value = head
    r = 2
    for item in events:
        ws.cell(r, 1).value = int(item.get("year") or BASE_YEAR)
        ws.cell(r, 2).value = int(item.get("month") or 0)
        ws.cell(r, 3).value = item.get("category") or ""
        ws.cell(r, 4).value = item.get("title") or ""
        ws.cell(r, 5).value = item.get("source") or "manual"
        ws.cell(r, 6).value = item.get("auto_key") or ""
        ws.cell(r, 7).value = 1 if item.get("suppressed") else 0
        r += 1


def save_workbook(path: Path, cells: list[dict] | None = None, events: list[dict] | None = None) -> int:
    """Одна запись в тот же xlsx: ячейки факта и/или лист событий."""
    wb = load_workbook(path, data_only=False)
    written = 0
    try:
        errors = validate_workbook(wb, path.name)
        if errors:
            raise ExcelStructureError("; ".join(errors[:8]))
        if cells:
            written = _write_fact_cells(wb, cells)
        if events is not None:
            _write_events_sheet(wb, events)
        tmp = path.with_name(path.name + ".tmp")
        wb.save(tmp)
    finally:
        wb.close()
    tmp.replace(path)
    return written


def read_events_sheet(path: Path | None = None) -> list[dict] | None:
    s = _seed()
    excel = path or s.find_excel()
    if not excel or not excel.exists():
        return None
    wb = load_workbook(excel, data_only=True, read_only=True)
    try:
        if EVENTS_SHEET not in wb.sheetnames:
            return None
        ws = wb[EVENTS_SHEET]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(v is None or str(v).strip() == "" for v in row[:4]):
                continue
            year = int(s._num(row[0]) or BASE_YEAR)
            month = int(s._num(row[1]) or 0)
            if month < 1 or month > 12:
                continue
            title = _label(row[3]) if len(row) > 3 else ""
            suppressed = int(s._num(row[6])) if len(row) > 6 else 0
            if not title and not suppressed:
                continue
            rows.append({
                "year": year,
                "month": month,
                "category": _label(row[2]) if len(row) > 2 else "",
                "title": title,
                "source": _label(row[4]) if len(row) > 4 else "manual",
                "auto_key": _label(row[5]) if len(row) > 5 else "",
                "suppressed": 1 if suppressed else 0,
            })
        return rows
    finally:
        wb.close()


def now_stamp() -> str:
    return datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
